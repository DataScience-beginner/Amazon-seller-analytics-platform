from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.modules.imports.workbook import validate_ooxml_archive

FORMULA_VERSION = "selleros.gst-india.amazon-gstr1.v1"
MAX_GST_ROWS_PER_SHEET = 100_000
MAX_GST_COLUMNS_PER_SHEET = 64
SHEETS = {
    "B2B": ("Taxable Value", "Cess Amount"),
    "B2B CN (cdnr)": ("Taxable Value", "Cess Amount"),
    "B2CL CN (cdnur)": ("Taxable Value", "Cess Amount"),
    "B2C Large": ("Taxable Value", "Cess Amount"),
    "B2C Small": ("Taxable Value", "Cess Amount"),
    "HSN Summary": (
        "Total Value",
        "Taxable Value",
        "Integrated Tax Amount",
        "Central Tax Amount",
        "State/UT Tax Amount",
        "Cess Amount",
    ),
}


@dataclass(frozen=True, slots=True)
class ParsedGstr1:
    sections: dict[str, list[dict[str, Any]]]
    totals: dict[str, str]
    row_counts: dict[str, int]
    configuration_checksum: str


def parse_amazon_gstr1(path: Path) -> ParsedGstr1:
    validate_ooxml_archive(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    missing = sorted(set(SHEETS) - set(workbook.sheetnames))
    if missing:
        raise ValueError(f"Missing required GST sheets: {', '.join(missing)}")
    sections: dict[str, list[dict[str, Any]]] = {}
    totals: dict[str, Decimal] = {}
    for sheet_name, money_headers in SHEETS.items():
        sheet = workbook[sheet_name]
        if sheet.max_row > MAX_GST_ROWS_PER_SHEET or sheet.max_column > MAX_GST_COLUMNS_PER_SHEET:
            raise ValueError(f"{sheet_name} exceeds the supported GST workbook dimensions")
        headers = [
            str(value).strip() if value is not None else ""
            for value in next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
        ]
        if not any(headers):
            raise ValueError(f"{sheet_name} has no header row")
        rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=5, values_only=True):
            if not any(value not in (None, "") for value in values):
                continue
            row = {
                header: _json_value(value)
                for header, value in zip(headers, values, strict=True)
                if header
            }
            rows.append(row)
            for header in money_headers:
                amount = _decimal(row.get(header))
                prefix = "hsn" if sheet_name == "HSN Summary" else "outward"
                total_key = f"{prefix}_{header.casefold().replace('/', '_').replace(' ', '_')}"
                totals[total_key] = totals.get(total_key, Decimal("0")) + amount
        sections[sheet_name] = rows
    manifest = json.dumps({"formula": FORMULA_VERSION, "sheets": SHEETS}, sort_keys=True)
    return ParsedGstr1(
        sections=sections,
        totals={key: str(value.quantize(Decimal("0.01"))) for key, value in totals.items()},
        row_counts={key: len(value) for key, value in sections.items()},
        configuration_checksum=hashlib.sha256(manifest.encode()).hexdigest(),
    )


def period_from_filename(filename: str) -> date | None:
    months = {
        name: index
        for index, name in enumerate(
            (
                "JANUARY",
                "FEBRUARY",
                "MARCH",
                "APRIL",
                "MAY",
                "JUNE",
                "JULY",
                "AUGUST",
                "SEPTEMBER",
                "OCTOBER",
                "NOVEMBER",
                "DECEMBER",
            ),
            start=1,
        )
    }
    parts = filename.upper().replace(".", "-").split("-")
    for index, part in enumerate(parts[:-1]):
        if part in months and parts[index + 1].isdigit():
            return date(int(parts[index + 1]), months[part], 1)
    return None


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError("A GST money field is not numeric") from error


def _json_value(value: Any) -> Any:
    if isinstance(value, str | int | bool) or value is None:
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(Decimal(str(value)))
    return str(value)
