from __future__ import annotations

import hashlib
import io
import math
import unicodedata
from collections.abc import Iterator, Mapping
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import BinaryIO, Protocol, cast
from xml.etree.ElementTree import ParseError
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, LargeZipFile, ZipFile, ZipInfo

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.modules.imports.domain import (
    CanonicalValueType,
    CellIssue,
    JsonScalar,
    MappingReport,
    SourceCell,
    WorkbookErrorCode,
    WorkbookInspection,
    WorkbookInspectionError,
    WorkbookRow,
    WorkbookSafetyLimits,
)
from app.modules.imports.mapping import map_headers
from app.modules.imports.normalization import normalize_header
from app.modules.imports.registry import AliasRegistry

WorkbookSource = Path | bytes

_REQUIRED_OOXML_MEMBERS = frozenset({"[content_types].xml", "xl/workbook.xml"})
_DANGEROUS_MEMBER_NAMES = frozenset({"xl/vbaproject.bin"})
_DANGEROUS_MEMBER_PREFIXES = (
    "xl/activex/",
    "xl/embeddings/",
    "xl/ctrlprops/",
    "customui/",
)
_XML_DECLARATIONS = (b"<!doctype", b"<!entity")
_XML_SCAN_CHUNK_BYTES = 64 * 1024


class _Worksheet(Protocol):
    title: str
    sheet_state: str
    max_row: int | None
    max_column: int | None

    def calculate_dimension(self, force: bool = False) -> str: ...

    def iter_rows(
        self,
        *,
        min_row: int | None = None,
        max_row: int | None = None,
        min_col: int | None = None,
        max_col: int | None = None,
        values_only: bool = False,
    ) -> Iterator[tuple[object, ...]]: ...


class _Workbook(Protocol):
    sheetnames: list[str]
    worksheets: list[_Worksheet]

    def __getitem__(self, key: str) -> _Worksheet: ...

    def close(self) -> None: ...


@dataclass(frozen=True, slots=True)
class _HeaderCandidate:
    sheet_name: str
    sheet_visible: bool
    row_number: int
    headers: tuple[str, ...]
    score: tuple[int, int, int, int, int]


def inspect_workbook(
    source: WorkbookSource,
    registry: AliasRegistry,
    *,
    explicit_mappings: Mapping[int, str | None] | None = None,
    limits: WorkbookSafetyLimits | None = None,
) -> WorkbookInspection:
    """Inspect an OOXML workbook and return a safe, typed mapping preview."""

    active_limits = limits or WorkbookSafetyLimits()
    checksum = _validate_ooxml_archive(source, active_limits)
    workbook = _load_workbook(source)
    try:
        _validate_workbook_dimensions(workbook, active_limits)
        candidate = _detect_header(workbook, registry, active_limits)
        worksheet = workbook[candidate.sheet_name]
        mapping = map_headers(
            candidate.headers,
            registry,
            explicit_mappings=explicit_mappings,
        )
        preview_rows = tuple(
            _iter_sheet_rows(
                worksheet,
                header_row=candidate.row_number,
                mapping=mapping,
                registry=registry,
                max_scanned_rows=active_limits.max_preview_scan_rows,
                max_data_rows=active_limits.preview_rows,
                truncate_when_full=True,
            )
        )
        if not preview_rows:
            raise WorkbookInspectionError(
                WorkbookErrorCode.empty_workbook,
                "The workbook contains headers but no product rows.",
                {"sheet": candidate.sheet_name, "header_row": candidate.row_number},
            )
        columns = tuple(entry.column for entry in mapping.columns if entry.column is not None)
        return WorkbookInspection(
            checksum_sha256=checksum,
            sheet_names=tuple(workbook.sheetnames),
            selected_sheet=candidate.sheet_name,
            header_row=candidate.row_number,
            columns=columns,
            mapping=mapping,
            preview_rows=preview_rows,
        )
    finally:
        workbook.close()


def iter_workbook_rows(
    source: WorkbookSource,
    inspection: WorkbookInspection,
    registry: AliasRegistry,
    *,
    limits: WorkbookSafetyLimits | None = None,
) -> Iterator[WorkbookRow]:
    """Yield JSON-safe typed rows after verifying the inspected file is unchanged."""

    active_limits = limits or WorkbookSafetyLimits()
    if (
        registry.registry_id != inspection.mapping.registry_id
        or registry.version != inspection.mapping.registry_version
    ):
        raise WorkbookInspectionError(
            WorkbookErrorCode.registry_changed,
            "The alias registry does not match the workbook mapping preview.",
            {
                "expected_registry_id": inspection.mapping.registry_id,
                "expected_registry_version": inspection.mapping.registry_version,
                "actual_registry_id": registry.registry_id,
                "actual_registry_version": registry.version,
            },
        )
    checksum = _validate_ooxml_archive(source, active_limits)
    if checksum != inspection.checksum_sha256:
        raise WorkbookInspectionError(
            WorkbookErrorCode.workbook_changed,
            "The workbook changed after its mapping preview was created.",
        )

    workbook = _load_workbook(source)
    try:
        _validate_workbook_dimensions(workbook, active_limits)
        if inspection.selected_sheet not in workbook.sheetnames:
            raise WorkbookInspectionError(
                WorkbookErrorCode.workbook_changed,
                "The selected worksheet no longer exists.",
            )
        worksheet = workbook[inspection.selected_sheet]
        observed_headers = _read_header_row(
            worksheet,
            inspection.header_row,
            active_limits.max_columns,
        )
        expected_headers = tuple(column.source_header for column in inspection.columns)
        if observed_headers != expected_headers:
            raise WorkbookInspectionError(
                WorkbookErrorCode.workbook_changed,
                "The workbook headers changed after its mapping preview was created.",
            )
        yield from _iter_sheet_rows(
            worksheet,
            header_row=inspection.header_row,
            mapping=inspection.mapping,
            registry=registry,
            max_scanned_rows=active_limits.max_data_scan_rows,
            max_data_rows=active_limits.max_data_rows,
            truncate_when_full=False,
        )
    finally:
        workbook.close()


def read_workbook_rows(
    source: WorkbookSource,
    inspection: WorkbookInspection,
    registry: AliasRegistry,
    *,
    limits: WorkbookSafetyLimits | None = None,
) -> tuple[WorkbookRow, ...]:
    """Materialize rows for bounded callers and tests."""

    return tuple(iter_workbook_rows(source, inspection, registry, limits=limits))


def _validate_ooxml_archive(source: WorkbookSource, limits: WorkbookSafetyLimits) -> str:
    archive_size = _source_size(source)
    if archive_size == 0:
        raise WorkbookInspectionError(
            WorkbookErrorCode.empty_workbook,
            "The uploaded workbook is empty.",
        )
    if archive_size > limits.max_archive_bytes:
        raise WorkbookInspectionError(
            WorkbookErrorCode.safety_limit_exceeded,
            "The workbook exceeds the configured compressed-size limit.",
            {"limit_bytes": limits.max_archive_bytes},
        )

    checksum = _source_checksum(source)
    try:
        with ZipFile(_zip_source(source), mode="r") as archive:
            members = archive.infolist()
            if len(members) > limits.max_archive_entries:
                raise WorkbookInspectionError(
                    WorkbookErrorCode.safety_limit_exceeded,
                    "The workbook contains too many archive members.",
                    {"limit": limits.max_archive_entries},
                )

            names: set[str] = set()
            total_uncompressed = 0
            for member in members:
                normalized_name = member.filename.replace("\\", "/")
                folded_name = normalized_name.casefold()
                path = PurePosixPath(normalized_name)
                if (
                    not normalized_name
                    or normalized_name.startswith("/")
                    or "\\" in member.filename
                    or ".." in path.parts
                    or len(normalized_name) > limits.max_filename_length
                ):
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.dangerous_archive,
                        "The workbook contains an unsafe archive path.",
                    )
                if folded_name in names:
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.dangerous_archive,
                        "The workbook contains duplicate archive paths.",
                    )
                names.add(folded_name)

                if member.flag_bits & 0x1:
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.dangerous_archive,
                        "Encrypted workbook members are not supported.",
                    )
                if member.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.dangerous_archive,
                        "The workbook uses an unsupported compression method.",
                    )
                if member.file_size > limits.max_member_uncompressed_bytes:
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.safety_limit_exceeded,
                        "A workbook member exceeds the configured size limit.",
                    )
                total_uncompressed += member.file_size
                if total_uncompressed > limits.max_total_uncompressed_bytes:
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.safety_limit_exceeded,
                        "The workbook exceeds the configured expanded-size limit.",
                    )
                if member.file_size > 0:
                    ratio = member.file_size / max(member.compress_size, 1)
                    if ratio > limits.max_compression_ratio:
                        raise WorkbookInspectionError(
                            WorkbookErrorCode.dangerous_archive,
                            "The workbook has an unsafe compression ratio.",
                        )

                if folded_name in _DANGEROUS_MEMBER_NAMES or folded_name.startswith(
                    _DANGEROUS_MEMBER_PREFIXES
                ):
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.dangerous_archive,
                        "Macro-enabled or embedded executable workbook content is not supported.",
                    )
                if (
                    not member.is_dir()
                    and folded_name.endswith((".xml", ".rels"))
                    and _contains_dangerous_xml_declaration(archive, member)
                ):
                    raise WorkbookInspectionError(
                        WorkbookErrorCode.dangerous_archive,
                        "Workbook XML declarations that can reference entities are not "
                        "supported.",
                    )

            if not _REQUIRED_OOXML_MEMBERS.issubset(names):
                raise WorkbookInspectionError(
                    WorkbookErrorCode.unsupported_structure,
                    "The file is not a supported .xlsx workbook.",
                )
    except WorkbookInspectionError:
        raise
    except (BadZipFile, LargeZipFile, OSError, RuntimeError) as exc:
        raise WorkbookInspectionError(
            WorkbookErrorCode.corrupt_workbook,
            "The workbook is corrupt or cannot be read.",
        ) from exc
    return checksum


def _contains_dangerous_xml_declaration(archive: ZipFile, member: ZipInfo) -> bool:
    """Scan a bounded OOXML member without assuming the declaration is near byte zero."""

    overlap = max(len(declaration) for declaration in _XML_DECLARATIONS) - 1
    carry = b""
    with archive.open(member) as member_file:
        while chunk := member_file.read(_XML_SCAN_CHUNK_BYTES):
            window = (carry + chunk).lower()
            if any(declaration in window for declaration in _XML_DECLARATIONS):
                return True
            carry = window[-overlap:]
    return False


def _load_workbook(source: WorkbookSource) -> _Workbook:
    try:
        workbook_source: Path | BinaryIO = (
            io.BytesIO(source) if isinstance(source, bytes) else source
        )
        return cast(
            _Workbook,
            load_workbook(
                workbook_source,
                read_only=True,
                data_only=True,
                keep_links=False,
            ),
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError) as exc:
        raise WorkbookInspectionError(
            WorkbookErrorCode.corrupt_workbook,
            "The workbook is corrupt or cannot be read.",
        ) from exc


def _validate_workbook_dimensions(workbook: _Workbook, limits: WorkbookSafetyLimits) -> None:
    if not workbook.worksheets:
        raise WorkbookInspectionError(
            WorkbookErrorCode.empty_workbook,
            "The workbook has no worksheets.",
        )
    if len(workbook.worksheets) > limits.max_worksheets:
        raise WorkbookInspectionError(
            WorkbookErrorCode.safety_limit_exceeded,
            "The workbook contains too many worksheets.",
            {"limit": limits.max_worksheets},
        )
    for worksheet in workbook.worksheets:
        if worksheet.max_column is None or worksheet.max_row is None:
            try:
                worksheet.calculate_dimension(force=True)
            except UnboundLocalError:
                # OpenPyXL leaves a truly empty, unsized worksheet without a final cell.
                # Header detection below will return the stable empty-workbook error.
                pass
            except (OSError, ParseError, ValueError) as exc:
                raise WorkbookInspectionError(
                    WorkbookErrorCode.corrupt_workbook,
                    "A worksheet has invalid dimension metadata and cannot be read safely.",
                    {"sheet": worksheet.title},
                ) from exc

        max_column = worksheet.max_column or 0
        max_row = worksheet.max_row or 0
        if max_column > limits.max_columns:
            raise WorkbookInspectionError(
                WorkbookErrorCode.safety_limit_exceeded,
                "A worksheet exceeds the configured column limit.",
                {"limit": limits.max_columns, "sheet": worksheet.title},
            )
        if max_row > limits.max_data_scan_rows + limits.max_header_scan_rows:
            raise WorkbookInspectionError(
                WorkbookErrorCode.safety_limit_exceeded,
                "A worksheet exceeds the configured row limit.",
                {"limit": limits.max_data_scan_rows, "sheet": worksheet.title},
            )


def _detect_header(
    workbook: _Workbook,
    registry: AliasRegistry,
    limits: WorkbookSafetyLimits,
) -> _HeaderCandidate:
    candidates: list[_HeaderCandidate] = []
    saw_value = False
    field_lookup = registry.fields_by_name

    for worksheet in workbook.worksheets:
        for row_number, raw_row in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=limits.max_header_scan_rows,
                max_col=limits.max_columns,
                values_only=True,
            ),
            start=1,
        ):
            if any(not _is_blank(value) for value in raw_row):
                saw_value = True
            headers = _trim_headers(raw_row)
            nonblank_count = sum(bool(header) for header in headers)
            if nonblank_count < limits.min_header_cells:
                continue

            candidate_sets = [
                registry.candidates_for(normalize_header(header)) for header in headers
            ]
            known_fields = {field for field_set in candidate_sets for field in field_set}
            if not known_fields:
                continue
            required_fields = {field for field in known_fields if field_lookup[field].required}
            ambiguous_columns = sum(len(field_set) > 1 for field_set in candidate_sets)
            visible = worksheet.sheet_state == "visible"
            score = (
                len(required_fields),
                len(known_fields),
                int(visible),
                -ambiguous_columns,
                nonblank_count,
            )
            candidates.append(
                _HeaderCandidate(
                    sheet_name=worksheet.title,
                    sheet_visible=visible,
                    row_number=row_number,
                    headers=headers,
                    score=score,
                )
            )

    if not candidates:
        if not saw_value:
            raise WorkbookInspectionError(
                WorkbookErrorCode.empty_workbook,
                "The workbook does not contain any data.",
            )
        raise WorkbookInspectionError(
            WorkbookErrorCode.unsupported_structure,
            "No supported Keepa header row could be detected.",
        )

    best_score = max(candidate.score for candidate in candidates)
    best = [candidate for candidate in candidates if candidate.score == best_score]
    if len(best) != 1:
        locations = ", ".join(
            f"{candidate.sheet_name}!{candidate.row_number}"
            for candidate in sorted(best, key=lambda item: (item.sheet_name, item.row_number))
        )
        raise WorkbookInspectionError(
            WorkbookErrorCode.ambiguous_header,
            "More than one equally likely Keepa header row was detected.",
            {"candidates": locations},
        )
    return best[0]


def _read_header_row(worksheet: _Worksheet, row_number: int, max_columns: int) -> tuple[str, ...]:
    rows = worksheet.iter_rows(
        min_row=row_number,
        max_row=row_number,
        max_col=max_columns,
        values_only=True,
    )
    try:
        raw_row = next(rows)
    except StopIteration as exc:
        raise WorkbookInspectionError(
            WorkbookErrorCode.workbook_changed,
            "The selected header row no longer exists.",
        ) from exc
    return _trim_headers(raw_row)


def _trim_headers(row: tuple[object, ...]) -> tuple[str, ...]:
    headers = [_header_text(value) for value in row]
    while headers and not headers[-1]:
        headers.pop()
    return tuple(headers)


def _header_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, date | datetime | time):
        return value.isoformat()
    return str(value).strip()


def _iter_sheet_rows(
    worksheet: _Worksheet,
    *,
    header_row: int,
    mapping: MappingReport,
    registry: AliasRegistry,
    max_scanned_rows: int,
    max_data_rows: int,
    truncate_when_full: bool,
) -> Iterator[WorkbookRow]:
    column_count = len(mapping.columns)
    if column_count == 0:
        return

    rows_seen = 0
    for offset, raw_row in enumerate(
        worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=header_row + max_scanned_rows,
            max_col=column_count,
            values_only=True,
        ),
        start=1,
    ):
        if all(_is_blank(value) for value in raw_row):
            continue
        rows_seen += 1
        if rows_seen > max_data_rows:
            if truncate_when_full:
                return
            raise WorkbookInspectionError(
                WorkbookErrorCode.safety_limit_exceeded,
                "The workbook contains more product rows than configured.",
                {"limit": max_data_rows},
            )
        yield _build_row(
            row_number=header_row + offset,
            raw_row=raw_row,
            mapping=mapping,
            registry=registry,
        )


def _build_row(
    *,
    row_number: int,
    raw_row: tuple[object, ...],
    mapping: MappingReport,
    registry: AliasRegistry,
) -> WorkbookRow:
    issues: list[CellIssue] = []
    source_values: list[SourceCell] = []
    raw_by_ordinal: dict[int, object] = {}
    for entry in mapping.columns:
        if entry.column is None:
            continue
        ordinal = entry.column.ordinal
        raw_value = raw_row[ordinal - 1] if ordinal <= len(raw_row) else None
        raw_by_ordinal[ordinal] = raw_value
        safe_value, issue_code = _json_safe_cell(raw_value)
        source_values.append(
            SourceCell(
                ordinal=ordinal,
                source_header=entry.column.source_header,
                value=safe_value,
            )
        )
        if issue_code is not None:
            issues.append(
                CellIssue(
                    code=issue_code,
                    message="The source value is not a finite JSON number.",
                    column_ordinal=ordinal,
                )
            )

    canonical_values: dict[str, JsonScalar] = {}
    for field_name, entry in mapping.resolved_by_field.items():
        if entry.column is None:
            continue
        field = registry.get_field(field_name)
        raw_value = raw_by_ordinal.get(entry.column.ordinal)
        value, issue = _coerce_canonical_value(raw_value, field.value_type)
        canonical_values[field_name] = value
        if issue is not None:
            issues.append(
                CellIssue(
                    code=issue,
                    message=f"The value could not be parsed as {field.value_type.value}.",
                    column_ordinal=entry.column.ordinal,
                    canonical_field=field_name,
                )
            )
        if field.required and _is_blank(value):
            issues.append(
                CellIssue(
                    code="missing_required_value",
                    message=f"Required field '{field_name}' is empty; no fallback is generated.",
                    column_ordinal=entry.column.ordinal,
                    canonical_field=field_name,
                )
            )

    return WorkbookRow(
        row_number=row_number,
        canonical_values=canonical_values,
        source_values=tuple(source_values),
        issues=tuple(issues),
    )


def _coerce_canonical_value(
    value: object,
    value_type: CanonicalValueType,
) -> tuple[JsonScalar, str | None]:
    if _is_blank(value):
        return None, None

    if value_type in (CanonicalValueType.text, CanonicalValueType.url):
        safe_value, issue = _json_safe_cell(value)
        return (str(safe_value).strip() if safe_value is not None else None), issue

    if value_type in (CanonicalValueType.decimal, CanonicalValueType.percentage):
        decimal_value = _parse_decimal(value)
        if decimal_value is None:
            return None, f"invalid_{value_type.value}"
        return format(decimal_value, "f"), None

    if value_type is CanonicalValueType.integer:
        decimal_value = _parse_decimal(value)
        if decimal_value is None or decimal_value != decimal_value.to_integral_value():
            return None, "invalid_integer"
        return int(decimal_value), None

    if value_type is CanonicalValueType.boolean:
        if isinstance(value, bool):
            return value, None
        if isinstance(value, int) and value in (0, 1):
            return bool(value), None
        normalized = str(value).strip().casefold()
        if normalized in {"yes", "true", "y", "1"}:
            return True, None
        if normalized in {"no", "false", "n", "0"}:
            return False, None
        return None, "invalid_boolean"

    if value_type is CanonicalValueType.date:
        if isinstance(value, datetime):
            return value.date().isoformat(), None
        if isinstance(value, date):
            return value.isoformat(), None
        return None, "invalid_date"

    if value_type is CanonicalValueType.datetime:
        if isinstance(value, datetime):
            return value.isoformat(), None
        if isinstance(value, date):
            return datetime.combine(value, time.min).isoformat(), None
        return None, "invalid_datetime"

    return None, "unsupported_value_type"


def _parse_decimal(value: object) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        return Decimal(str(value))

    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    cleaned = "".join(
        character for character in text if not unicodedata.category(character).startswith("Sc")
    )
    cleaned = cleaned.replace(",", "").replace("%", "").replace(" ", "")
    if negative:
        cleaned = f"-{cleaned}"
    try:
        parsed = Decimal(cleaned)
    except InvalidOperation:
        return None
    return parsed if parsed.is_finite() else None


def _json_safe_cell(value: object) -> tuple[JsonScalar, str | None]:
    if value is None or isinstance(value, str | bool | int):
        return cast(JsonScalar, value), None
    if isinstance(value, float):
        if math.isfinite(value):
            return value, None
        if math.isnan(value):
            return "NaN", "non_finite_number"
        return ("Infinity" if value > 0 else "-Infinity"), "non_finite_number"
    if isinstance(value, Decimal):
        if value.is_finite():
            return format(value, "f"), None
        return str(value), "non_finite_number"
    if isinstance(value, datetime | date | time):
        return value.isoformat(), None
    return str(value), None


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _source_size(source: WorkbookSource) -> int:
    return len(source) if isinstance(source, bytes) else source.stat().st_size


def _source_checksum(source: WorkbookSource) -> str:
    digest = hashlib.sha256()
    if isinstance(source, bytes):
        digest.update(source)
        return digest.hexdigest()
    with source.open("rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _zip_source(source: WorkbookSource) -> Path | BinaryIO:
    return io.BytesIO(source) if isinstance(source, bytes) else source
