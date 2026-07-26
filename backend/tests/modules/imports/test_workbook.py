from __future__ import annotations

import io
import re
from dataclasses import replace
from datetime import date, datetime
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.modules.imports import (
    AliasRegistry,
    MappingClassification,
    WorkbookErrorCode,
    WorkbookInspectionError,
    WorkbookSafetyLimits,
    inspect_workbook,
    read_workbook_rows,
)


def _assert_error(
    source: bytes,
    registry: AliasRegistry,
    code: WorkbookErrorCode,
    *,
    limits: WorkbookSafetyLimits | None = None,
) -> WorkbookInspectionError:
    with pytest.raises(WorkbookInspectionError) as caught:
        inspect_workbook(source, registry, limits=limits)
    assert caught.value.code is code
    assert str(caught.value) == caught.value.user_message
    return caught.value


def _typed_registry() -> AliasRegistry:
    fields = [
        ("asin", True, "text", "ASIN"),
        ("quantity", False, "integer", "Quantity"),
        ("price", False, "decimal", "Price"),
        ("rate", False, "percentage", "Rate"),
        ("is_fba", False, "boolean", "Is FBA"),
        ("observed_date", False, "date", "Observed Date"),
        ("observed_at", False, "datetime", "Observed At"),
        ("listing_url", False, "url", "Listing URL"),
    ]
    return AliasRegistry.from_object(
        {
            "schema_version": 1,
            "registry_id": "typed_test",
            "version": "1",
            "fields": [
                {
                    "canonical_field": name,
                    "required": required,
                    "value_type": value_type,
                    "aliases": [alias],
                    "description": f"Test field {name}.",
                }
                for name, required, value_type, alias in fields
            ],
        }
    )


def test_detects_sheet_and_header_after_preamble(registry: AliasRegistry) -> None:
    workbook = Workbook()
    cover = workbook.active
    assert cover is not None
    cover.title = "Read me"
    cover.append(["Keepa export", "Generated for testing"])
    products = workbook.create_sheet("Products")
    products.append(["Synthetic export", None, None])
    products.append([None, None, None])
    products.append(["ASIN", "Product Name", "Mystery Metric"])
    products.append(["B000000001", "Synthetic widget", 7])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    inspection = inspect_workbook(output.getvalue(), registry)

    assert inspection.selected_sheet == "Products"
    assert inspection.header_row == 3
    assert inspection.sheet_names == ("Read me", "Products")
    assert [column.source_header for column in inspection.columns] == [
        "ASIN",
        "Product Name",
        "Mystery Metric",
    ]
    assert len(inspection.checksum_sha256) == 64
    assert inspection.preview_rows[0].row_number == 4


def test_recalculates_missing_worksheet_dimension_metadata(
    registry: AliasRegistry,
    make_workbook_bytes: Any,
) -> None:
    source = make_workbook_bytes([["ASIN", "Title"], ["B000000001", "Synthetic unsized worksheet"]])
    output = io.BytesIO()
    dimension_pattern = re.compile(rb"<dimension\s+[^>]*/>")
    with (
        ZipFile(io.BytesIO(source), "r") as original,
        ZipFile(output, "w", compression=ZIP_DEFLATED) as modified,
    ):
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                content, replacements = dimension_pattern.subn(b"", content, count=1)
                assert replacements == 1
            modified.writestr(member, content)

    inspection = inspect_workbook(output.getvalue(), registry)

    assert inspection.selected_sheet == "Products"
    assert inspection.preview_rows[0].canonical_values["asin"] == "B000000001"


def test_equally_likely_headers_are_reported_as_ambiguous(registry: AliasRegistry) -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "First"
    first.append(["ASIN", "Title"])
    first.append(["B000000001", "One"])
    second = workbook.create_sheet("Second")
    second.append(["ASIN", "Title"])
    second.append(["B000000002", "Two"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    error = _assert_error(output.getvalue(), registry, WorkbookErrorCode.ambiguous_header)
    assert error.details["candidates"] == "First!1, Second!1"


def test_unknown_columns_and_values_are_preserved(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    source = make_workbook_bytes(
        [
            ["ASIN", "Title", "Future Keepa Metric"],
            ["B000000001", "Synthetic product", 123.5],
        ]
    )

    inspection = inspect_workbook(source, registry)
    row = inspection.preview_rows[0]

    assert inspection.mapping.columns[2].classification is MappingClassification.unknown
    assert row.canonical_values == {
        "asin": "B000000001",
        "title": "Synthetic product",
    }
    assert [(cell.source_header, cell.value) for cell in row.source_values] == [
        ("ASIN", "B000000001"),
        ("Title", "Synthetic product"),
        ("Future Keepa Metric", 123.5),
    ]


def test_formula_cells_return_cached_values_without_evaluating_expressions(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    source = make_workbook_bytes(
        [
            ["ASIN", "Title", "Future Keepa Metric"],
            ["B000000001", "=1+1", '=HYPERLINK("https://invalid.example", "click")'],
        ]
    )
    output = io.BytesIO()
    with (
        ZipFile(io.BytesIO(source), "r") as original,
        ZipFile(output, "w", compression=ZIP_DEFLATED) as modified,
    ):
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                formula_without_cache = b"<f>1+1</f><v />"
                assert formula_without_cache in content
                content = content.replace(
                    formula_without_cache,
                    b"<f>1+1</f><v>999</v>",
                    1,
                )
            modified.writestr(member, content)

    inspection = inspect_workbook(output.getvalue(), registry)
    row = inspection.preview_rows[0]

    assert row.canonical_values["title"] == "999"
    assert row.source_values[1].value == 999
    assert row.source_values[2].value is None
    assert all(
        not (isinstance(cell.value, str) and cell.value.startswith("="))
        for cell in row.source_values
    )


def test_rows_are_typed_without_binary_float_money(
    make_workbook_bytes: Any,
) -> None:
    registry = _typed_registry()
    source = make_workbook_bytes(
        [
            [
                "ASIN",
                "Quantity",
                "Price",
                "Rate",
                "Is FBA",
                "Observed Date",
                "Observed At",
                "Listing URL",
            ],
            [
                " B000000001 ",
                42,
                "$1,234.50",
                "12.5%",
                "YES",
                date(2026, 1, 2),
                datetime(2026, 1, 2, 3, 4, 5),
                " https://example.invalid/product ",
            ],
        ]
    )

    inspection = inspect_workbook(source, registry)
    row = read_workbook_rows(source, inspection, registry)[0]

    assert row.canonical_values == {
        "asin": "B000000001",
        "quantity": 42,
        "price": "1234.50",
        "rate": "12.5",
        "is_fba": True,
        "observed_date": "2026-01-02",
        "observed_at": "2026-01-02T03:04:05",
        "listing_url": "https://example.invalid/product",
    }
    assert row.issues == ()


def test_malformed_cells_create_issues_and_keep_source_values(
    make_workbook_bytes: Any,
) -> None:
    registry = _typed_registry()
    source = make_workbook_bytes(
        [
            ["ASIN", "Quantity", "Price", "Is FBA", "Observed Date"],
            [None, "1.5", "not money", "perhaps", "2026-01-02"],
        ]
    )

    inspection = inspect_workbook(source, registry)
    row = inspection.preview_rows[0]

    assert row.canonical_values == {
        "asin": None,
        "quantity": None,
        "price": None,
        "is_fba": None,
        "observed_date": None,
    }
    assert {issue.code for issue in row.issues} == {
        "missing_required_value",
        "invalid_integer",
        "invalid_decimal",
        "invalid_boolean",
        "invalid_date",
    }
    assert [cell.value for cell in row.source_values] == [
        None,
        "1.5",
        "not money",
        "perhaps",
        "2026-01-02",
    ]


def test_blank_rows_are_skipped_without_losing_excel_row_numbers(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    source = make_workbook_bytes(
        [
            ["ASIN", "Title"],
            ["B000000001", "First"],
            [None, None],
            ["B000000002", "Second"],
        ]
    )

    inspection = inspect_workbook(source, registry)
    rows = read_workbook_rows(source, inspection, registry)

    assert [row.row_number for row in rows] == [2, 4]
    assert [row.canonical_values["asin"] for row in rows] == [
        "B000000001",
        "B000000002",
    ]


@pytest.mark.parametrize(
    ("source", "code"),
    [
        (b"", WorkbookErrorCode.empty_workbook),
        (b"not a zip archive", WorkbookErrorCode.corrupt_workbook),
    ],
)
def test_empty_and_malformed_uploads_have_stable_errors(
    source: bytes, code: WorkbookErrorCode, registry: AliasRegistry
) -> None:
    _assert_error(source, registry, code)


def test_unsupported_and_header_only_workbooks_are_rejected(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    unsupported = make_workbook_bytes([["Unrelated column", "Another column"], ["one", "two"]])
    header_only = make_workbook_bytes([["ASIN", "Title"]])

    _assert_error(unsupported, registry, WorkbookErrorCode.unsupported_structure)
    _assert_error(header_only, registry, WorkbookErrorCode.empty_workbook)


def test_zip_without_required_ooxml_members_is_unsupported(registry: AliasRegistry) -> None:
    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("notes.txt", "not a workbook")

    _assert_error(output.getvalue(), registry, WorkbookErrorCode.unsupported_structure)


def test_checksum_prevents_parsing_a_changed_workbook(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    first = make_workbook_bytes([["ASIN", "Title"], ["B000000001", "First"]])
    second = make_workbook_bytes([["ASIN", "Title"], ["B000000002", "Second"]])
    inspection = inspect_workbook(first, registry)

    with pytest.raises(WorkbookInspectionError) as caught:
        read_workbook_rows(second, inspection, registry)
    assert caught.value.code is WorkbookErrorCode.workbook_changed


def test_registry_version_is_locked_between_preview_and_row_parsing(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    source = make_workbook_bytes([["ASIN", "Title"], ["B000000001", "First"]])
    inspection = inspect_workbook(source, registry)
    changed_registry = AliasRegistry(
        registry_id=registry.registry_id,
        version="2.0.0",
        fields=registry.fields,
    )

    with pytest.raises(WorkbookInspectionError) as caught:
        read_workbook_rows(source, inspection, changed_registry)
    assert caught.value.code is WorkbookErrorCode.registry_changed
    assert caught.value.details == {
        "expected_registry_id": "keepa",
        "expected_registry_version": "1.0.0",
        "actual_registry_id": "keepa",
        "actual_registry_version": "2.0.0",
    }


def test_configured_archive_and_sheet_limits_are_enforced(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    source = make_workbook_bytes([["ASIN", "Title", "Brand"], ["B000000001", "Product", "Brand"]])
    defaults = WorkbookSafetyLimits()

    cases = [
        replace(defaults, max_archive_bytes=len(source) - 1),
        replace(defaults, max_archive_entries=1),
        replace(defaults, max_member_uncompressed_bytes=10),
        replace(defaults, max_total_uncompressed_bytes=10),
        replace(defaults, max_compression_ratio=1.0),
        replace(defaults, max_columns=2),
    ]
    for limits in cases:
        _assert_error(
            source,
            registry,
            WorkbookErrorCode.safety_limit_exceeded
            if limits.max_compression_ratio != 1.0
            else WorkbookErrorCode.dangerous_archive,
            limits=limits,
        )


@pytest.mark.parametrize(
    ("member_name", "content"),
    [
        ("../payload.bin", b"payload"),
        ("xl/vbaProject.bin", b"macro"),
        ("xl/activeX/control.bin", b"control"),
        ("XL/WORKBOOK.XML", b"duplicate using different case"),
    ],
)
def test_dangerous_archive_members_are_rejected(
    member_name: str,
    content: bytes,
    registry: AliasRegistry,
    make_workbook_bytes: Any,
    add_archive_member: Any,
) -> None:
    source = make_workbook_bytes([["ASIN", "Title"], ["B000000001", "Product"]])
    dangerous = add_archive_member(source, member_name, content)

    _assert_error(dangerous, registry, WorkbookErrorCode.dangerous_archive)


def test_entity_declarations_in_workbook_xml_are_rejected(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    source = make_workbook_bytes([["ASIN", "Title"], ["B000000001", "Product"]])
    output = io.BytesIO()
    with (
        ZipFile(io.BytesIO(source), "r") as original,
        ZipFile(output, "w", compression=ZIP_DEFLATED) as modified,
    ):
        for member in original.infolist():
            content = original.read(member.filename)
            if member.filename == "xl/workbook.xml":
                content = b" " * 9000 + b"<!DOCTYPE workbook>" + content
            modified.writestr(member, content)

    _assert_error(output.getvalue(), registry, WorkbookErrorCode.dangerous_archive)


def test_too_many_data_rows_are_rejected_instead_of_truncated(
    registry: AliasRegistry, make_workbook_bytes: Any
) -> None:
    source = make_workbook_bytes(
        [
            ["ASIN", "Title"],
            ["B000000001", "First"],
            ["B000000002", "Second"],
        ]
    )
    inspection = inspect_workbook(source, registry)
    limits = replace(WorkbookSafetyLimits(), max_data_rows=1)

    with pytest.raises(WorkbookInspectionError) as caught:
        read_workbook_rows(source, inspection, registry, limits=limits)
    assert caught.value.code is WorkbookErrorCode.safety_limit_exceeded
